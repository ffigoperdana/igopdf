import os
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

try:
    from .docx_convert import (
        is_nota_dinas,
        is_nota_riil,
        is_rincian_biaya_perjalanan_dinas,
        repair_editable_docx,
        restore_missing_spaces,
    )
except ImportError:
    from docx_convert import (
        is_nota_dinas,
        is_nota_riil,
        is_rincian_biaya_perjalanan_dinas,
        repair_editable_docx,
        restore_missing_spaces,
    )


class DocxConvertTextRepairTest(unittest.TestCase):
    def test_detects_nota_dinas_without_matching_generic_documents(self):
        self.assertTrue(
            is_nota_dinas(
                ["KEMENTERIAN KEUANGAN\nNOTA DINAS\nNOMOR ND-63/BPDP.100/2026"]
            )
        )
        self.assertFalse(
            is_nota_dinas(["SURAT UNDANGAN\nNOMOR UND-63/BPDP.100/2026"])
        )

    def test_detects_nota_riil_expense_form_without_matching_generic_tables(self):
        self.assertTrue(
            is_nota_riil(
                [
                    "DAFTAR PENGELUARAN RIIL\n"
                    "No Uraian Jumlah\n"
                    "Transportasi Dalam Kota"
                ]
            )
        )
        self.assertFalse(
            is_nota_riil(["No Uraian Jumlah\nRekap biaya perjalanan"])
        )

    def test_detects_rincian_biaya_form_without_matching_generic_tables(self):
        self.assertTrue(
            is_rincian_biaya_perjalanan_dinas(
                [
                    "RINCIAN BIAYA PERJALANAN DINAS\n"
                    "No Perincian Biaya Jumlah Keterangan"
                ]
            )
        )
        self.assertTrue(
            is_rincian_biaya_perjalanan_dinas(
                [
                    "RINCIANBIAYAPERJALANANDINAS\n"
                    "No PerincianBiaya Jumlah Keterangan"
                ]
            )
        )
        self.assertFalse(
            is_rincian_biaya_perjalanan_dinas(
                ["Rincian biaya rapat\nNo Uraian Jumlah"]
            )
        )

    def test_restores_realistic_nd_prose_from_pdf_whitespace(self):
        converted = (
            "Sehubungandenganpelaksanaanmonitoringdananalisisisumediamassadanmediasosial "
            "terkait APBN."
        )
        source = (
            "Sehubungan dengan pelaksanaan monitoring dan analisis isu media massa dan "
            "media sosial terkait APBN."
        )
        self.assertEqual(
            restore_missing_spaces(converted, [source]),
            source,
        )

    def test_restores_table_text_but_keeps_compound_hyphens(self):
        converted = (
            "Daftar ketentuan yang ditetapkanolehmasing-masingPimpinanEselonI."
        )
        source = (
            "Daftar ketentuan yang ditetapkan oleh masing-masing Pimpinan Eselon I."
        )
        self.assertEqual(
            restore_missing_spaces(converted, [source]),
            source,
        )

    def test_restores_table_cell_when_pdf_reading_order_interleaves_columns(self):
        converted = "DaftarsistemyangdigunakanolehUnitEselonI"
        full_page = (
            "Daftar sistem\n"
            "Divisi KLI\n"
            "yang digunakan oleh\n"
            "Unit Eselon I"
        )
        source_block = "Daftar sistem yang digunakan oleh Unit Eselon I"
        self.assertEqual(
            restore_missing_spaces(
                converted, [full_page, source_block], min_compact_chars=6
            ),
            "Daftar sistem yang digunakan oleh Unit Eselon I",
        )

    def test_unicode_casefold_expansion_keeps_reference_offsets_aligned(self):
        self.assertEqual(
            restore_missing_spaces(
                "VorgangTest",
                ["Maß Vorgang Test"],
                min_compact_chars=6,
            ),
            "Vorgang Test",
        )

    def test_leaves_ambiguous_spacing_and_identifiers_untouched(self):
        self.assertEqual(
            restore_missing_spaces("ABCD", ["AB CD", "A BCD"], min_compact_chars=4),
            "ABCD",
        )
        identifier_text = (
            "ABC123/2026 https://contoh.go.id user@example.go.id"
        )
        self.assertEqual(
            restore_missing_spaces(identifier_text, [identifier_text]),
            identifier_text,
        )
        self.assertEqual(
            restore_missing_spaces(
                "ABC123/2026",
                ["ABC 123/2026"],
                min_compact_chars=6,
            ),
            "ABC123/2026",
        )
        self.assertEqual(
            restore_missing_spaces(
                "https://contoh.go.id/path",
                ["https://contoh.go.id/ path"],
                min_compact_chars=6,
            ),
            "https://contoh.go.id/path",
        )
        self.assertEqual(
            restore_missing_spaces("ABCD", ["AB CD"], min_compact_chars=6),
            "ABCD",
        )
        self.assertEqual(
            restore_missing_spaces(
                "turunanatasPermenpan6Tahun2022.",
                ["turunan atas Permenpan 6 Tahun 2022."],
                min_compact_chars=6,
            ),
            "turunan atas Permenpan 6 Tahun 2022.",
        )
        self.assertEqual(
            restore_missing_spaces(
                "PimpinanEselonIpadaperiodeJanuari2026",
                [
                    "PimpinanEselonIpadaperiodeJanuari2026",
                    "Pimpinan Eselon I pada periode Januari 2026",
                ],
                min_compact_chars=6,
                prefer_richest_spacing=True,
            ),
            "Pimpinan Eselon I pada periode Januari 2026",
        )
        self.assertEqual(
            restore_missing_spaces(
                "ditetapkanolehmasing-masingPimpinanEselonIpadaperiodeJanuari2026",
                [
                    "ditetapkan oleh masing-masing Pimpinan Eselon I pada "
                    "periode Januari 2026"
                ],
                min_compact_chars=6,
                prefer_richest_spacing=True,
            ),
            "ditetapkan oleh masing-masing Pimpinan Eselon I pada periode "
            "Januari 2026",
        )

    def test_repairs_body_and_table_runs_without_losing_formatting(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "sample.docx")
            document = Document()
            paragraph = document.add_paragraph()
            bold_run = paragraph.add_run("Sehubungandengan")
            bold_run.bold = True
            paragraph.add_run("\tpelaksanaan")
            table = document.add_table(rows=2, cols=1)
            table.cell(0, 0).text = "Datayangmemuatidentitaspegawai"
            table.cell(1, 0).text = "Laporanyangdapatdisalin"
            document.save(path)

            restored = repair_editable_docx(
                path,
                [
                    "Sehubungan dengan pelaksanaan",
                    "Data yang memuat identitas pegawai",
                    "Laporan yang dapat disalin",
                ],
                nota_dinas=True,
            )

            repaired = Document(path)
            self.assertEqual(
                repaired.paragraphs[0].text,
                "Sehubungan dengan\tpelaksanaan",
            )
            self.assertTrue(repaired.paragraphs[0].runs[0].bold)
            self.assertEqual(
                repaired.tables[0].cell(0, 0).text,
                "Data yang memuat identitas pegawai",
            )
            self.assertEqual(
                repaired.tables[0].cell(1, 0).text,
                "Laporan yang dapat disalin",
            )
            self.assertEqual(restored, 8)

    def test_normalizes_nd_continuation_tabs_and_symbol_bullets(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "layout.docx")
            document = Document()
            numbered = document.add_paragraph()
            numbered.add_run("1. ")
            numbered.add_run("\t")
            numbered.add_run("Judul dan kalimat pertama ")
            numbered.add_run("\t")
            numbered.add_run("dilanjutkan secara alami.")
            numbered.paragraph_format.left_indent = Inches(0.1)
            numbered.paragraph_format.tab_stops.add_tab_stop(Inches(0.4))
            bullet = document.add_paragraph()
            symbol = bullet.add_run("\uf0b7")
            symbol.font.name = "SymbolMT"
            body = bullet.add_run("Isi butir yang portabel.")
            body.font.name = "Arial"
            heading = document.add_paragraph()
            heading.add_run("4.")
            heading.add_run("Aktivitas Media Sosial")
            document.save(path)

            restored = repair_editable_docx(
                path,
                [
                    "1. Judul dan kalimat pertama dilanjutkan secara alami.",
                    "\u2022 Isi butir yang portabel.",
                    "4. Aktivitas Media Sosial",
                ],
                nota_dinas=True,
            )

            repaired = Document(path)
            self.assertEqual(
                repaired.paragraphs[0].text,
                "1. \tJudul dan kalimat pertama dilanjutkan secara alami.",
            )
            self.assertEqual(
                repaired.paragraphs[1].text,
                "\u2022 Isi butir yang portabel.",
            )
            self.assertEqual(repaired.paragraphs[1].runs[0].font.name, "Arial")
            self.assertEqual(
                repaired.paragraphs[2].text,
                "4. Aktivitas Media Sosial",
            )
            self.assertEqual(
                repaired.paragraphs[0].paragraph_format.left_indent,
                Inches(0.4),
            )
            self.assertEqual(
                repaired.paragraphs[0].paragraph_format.first_line_indent,
                Inches(-0.3),
            )
            self.assertEqual(restored, 1)

    def test_rebuilds_merged_nd_tembusan_table_from_source_lines(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "tembusan.docx")
            document = Document()
            document.add_paragraph("Tembusan:")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "1.\n2.\n3."
            table.cell(0, 1).text = (
                "Penerima pertama Penerima kedua\nPenerima ketiga"
            )
            document.save(path)

            repair_editable_docx(
                path,
                [
                    "Tembusan:\n"
                    "1. Penerima pertama\n"
                    "2. Penerima kedua\n"
                    "3. Penerima ketiga"
                ],
                nota_dinas=True,
            )

            repaired = Document(path)
            self.assertEqual(len(repaired.tables[0].rows), 3)
            self.assertTrue(
                all(row.height is None for row in repaired.tables[0].rows)
            )
            self.assertEqual(
                [
                    [cell.text for cell in row.cells]
                    for row in repaired.tables[0].rows
                ],
                [
                    ["1.", "Penerima pertama"],
                    ["2.", "Penerima kedua"],
                    ["3.", "Penerima ketiga"],
                ],
            )

    def test_restores_visible_auto_colored_grid_for_nota_riil_expense_table(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "nota-riil.docx")
            document = Document()
            outer = document.add_table(rows=1, cols=1)
            table = outer.cell(0, 0).add_table(rows=3, cols=5)
            table.cell(0, 0).text = "No"
            table.cell(0, 1).merge(table.cell(0, 3)).text = "Uraian"
            table.cell(0, 4).text = "Jumlah"
            table.cell(1, 0).text = "1."
            table.cell(1, 1).merge(table.cell(1, 3)).text = "Transportasi Dalam Kota (PP)"
            table.cell(1, 4).text = "Rp. 170,000"
            table.cell(2, 0).text = "Jumlah"
            table.cell(2, 1).text = "1 x"
            table.cell(2, 2).text = "170,000.0 x"
            table.cell(2, 3).text = "100%"
            table.cell(2, 4).text = "Rp 170,000"
            document.save(path)

            repaired = repair_editable_docx(
                path,
                ["DAFTAR PENGELUARAN RIIL\nNo Uraian Jumlah"],
                nota_riil=True,
            )

            result = Document(path)
            expense_table = result.tables[0].cell(0, 0).tables[0]
            table_borders = expense_table._tbl.tblPr.first_child_found_in(
                "w:tblBorders"
            )
            self.assertEqual(repaired, 0)
            self.assertEqual(
                [cell.text for cell in expense_table.rows[1].cells],
                [
                    "1.",
                    "Transportasi Dalam Kota (PP)\n1 x 170,000.0 x 100%",
                    "Transportasi Dalam Kota (PP)\n1 x 170,000.0 x 100%",
                    "Transportasi Dalam Kota (PP)\n1 x 170,000.0 x 100%",
                    "Rp. 170,000",
                ],
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[2].cells],
                ["Jumlah", "Jumlah", "Jumlah", "Jumlah", "Rp 170,000"],
            )
            item_height = expense_table.rows[1]._tr.trPr.find(qn("w:trHeight"))
            self.assertEqual(item_height.get(qn("w:val")), "720")
            self.assertEqual(item_height.get(qn("w:hRule")), "atLeast")
            for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
                border = table_borders.find(qn(f"w:{edge}"))
                self.assertIsNotNone(border)
                self.assertEqual(border.get(qn("w:val")), "single")
                self.assertEqual(border.get(qn("w:color")), "auto")

    def test_repairs_multi_item_nota_riil_table_without_touching_prior_items(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "nota-riil-multi.docx")
            document = Document()
            outer = document.add_table(rows=1, cols=1)
            table = outer.cell(0, 0).add_table(rows=4, cols=5)
            table.cell(0, 0).text = "No"
            table.cell(0, 1).merge(table.cell(0, 3)).text = "Uraian"
            table.cell(0, 4).text = "Jumlah"
            table.cell(1, 0).text = "1."
            table.cell(1, 1).merge(table.cell(1, 3)).text = "Taksi bandara"
            table.cell(1, 4).text = "Rp. 120,000"
            table.cell(2, 0).text = "2."
            table.cell(2, 1).merge(table.cell(2, 3)).text = "Transportasi Dalam Kota (PP)"
            table.cell(2, 4).text = "Rp. 200,000"
            table.cell(3, 0).text = "Jumlah"
            table.cell(3, 1).text = "2 x"
            table.cell(3, 2).text = "100,000 x"
            table.cell(3, 3).text = "100%"
            table.cell(3, 4).text = "Rp 320,000"
            document.save(path)

            repair_editable_docx(
                path,
                ["DAFTAR PENGELUARAN RIIL\nNo Uraian Jumlah"],
                nota_riil=True,
            )

            result = Document(path)
            expense_table = result.tables[0].cell(0, 0).tables[0]
            self.assertEqual(
                expense_table.rows[1].cells[1].text,
                "Taksi bandara",
            )
            self.assertEqual(
                expense_table.rows[2].cells[1].text,
                "Transportasi Dalam Kota (PP)\n2 x 100,000 x 100%",
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[3].cells],
                ["Jumlah", "Jumlah", "Jumlah", "Jumlah", "Rp 320,000"],
            )
            second_item_height = expense_table.rows[2]._tr.trPr.find(
                qn("w:trHeight")
            )
            self.assertEqual(second_item_height.get(qn("w:val")), "720")
            self.assertEqual(second_item_height.get(qn("w:hRule")), "atLeast")

    def test_rebuilds_tab_collapsed_nota_riil_table_as_editable_grid(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "nota-riil-collapsed.docx")
            document = Document()
            outer = document.add_table(rows=1, cols=1)
            collapsed = outer.cell(0, 0).add_table(rows=1, cols=1)
            collapsed.cell(0, 0).text = (
                "No\tUraian\tJumlah\n"
                "1. Penginapan\tRp. 500,000\n"
                "2 malam x\t250,000.0 x\t100%\n"
                "2. Transportasi\tRp. 200,000\n"
                "1 x\t200,000.0 x\t100%\n"
                "Jumlah\tRp 700,000"
            )
            document.save(path)

            repair_editable_docx(
                path,
                ["DAFTAR PENGELUARAN RIIL\nNo Uraian Jumlah"],
                nota_riil=True,
            )

            result = Document(path)
            candidates = []

            def collect(table):
                candidates.append(table)
                seen = set()
                for row in table.rows:
                    for cell in row.cells:
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        for nested in cell.tables:
                            collect(nested)

            for table in result.tables:
                collect(table)
            expense_table = next(
                table
                for table in candidates
                if len(table.rows) == 4
                and len(table.columns) == 3
                and table.cell(0, 1).text == "Uraian"
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[1].cells],
                ["1.", "Penginapan\n2 malam x 250,000.0 x 100%", "Rp. 500,000"],
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[2].cells],
                ["2.", "Transportasi\n1 x 200,000.0 x 100%", "Rp. 200,000"],
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[3].cells],
                ["Jumlah", "Jumlah", "Rp 700,000"],
            )
            table_borders = expense_table._tbl.tblPr.first_child_found_in(
                "w:tblBorders"
            )
            self.assertEqual(
                table_borders.find(qn("w:insideV")).get(qn("w:color")), "auto"
            )

    def test_promotes_nota_riil_grid_out_of_exact_height_pdf2docx_wrappers(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "nota-riil-clipped.docx")
            document = Document()
            outer = document.add_table(rows=1, cols=1)
            middle = outer.cell(0, 0).add_table(rows=1, cols=1)
            for row, height_value in ((outer.rows[0], 1688), (middle.rows[0], 1538)):
                height = OxmlElement("w:trHeight")
                height.set(qn("w:val"), str(height_value))
                height.set(qn("w:hRule"), "exact")
                row._tr.get_or_add_trPr().append(height)

            table = middle.cell(0, 0).add_table(rows=3, cols=5)
            table.cell(0, 0).text = "No"
            table.cell(0, 1).merge(table.cell(0, 3)).text = "Uraian"
            table.cell(0, 4).text = "Jumlah"
            table.cell(1, 0).text = "1."
            table.cell(1, 1).merge(table.cell(1, 3)).text = (
                "Transportasi Dalam Kota (PP)\n1 x 170,000.0 x 100%"
            )
            table.cell(1, 4).text = "Rp. 170,000"
            table.cell(2, 0).merge(table.cell(2, 3)).text = "Jumlah"
            table.cell(2, 4).text = "Rp 170,000"
            document.save(path)

            repair_editable_docx(
                path,
                ["DAFTAR PENGELUARAN RIIL\nNo Uraian Jumlah"],
                nota_riil=True,
            )

            result = Document(path)
            self.assertEqual(len(result.tables), 1)
            expense_table = result.tables[0]
            self.assertEqual((len(expense_table.rows), len(expense_table.columns)), (3, 3))
            self.assertEqual(
                [cell.text for cell in expense_table.rows[0].cells],
                ["No", "Uraian", "Jumlah"],
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[1].cells],
                [
                    "1.",
                    "Transportasi Dalam Kota (PP)\n1 x 170,000.0 x 100%",
                    "Rp. 170,000",
                ],
            )
            self.assertEqual(
                [cell.text for cell in expense_table.rows[2].cells],
                ["Jumlah", "Jumlah", "Rp 170,000"],
            )
            self.assertFalse(expense_table.cell(0, 0).tables)
            header_shading = expense_table.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
            self.assertEqual(header_shading.get(qn("w:fill")), "808080")
            item_height = expense_table.rows[1]._tr.trPr.find(qn("w:trHeight"))
            self.assertEqual(item_height.get(qn("w:hRule")), "atLeast")

    def test_restores_light_gray_cells_in_rincian_biaya_form(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "rincian-biaya.docx")
            document = Document()
            total_table = document.add_table(rows=2, cols=3)
            total_table.cell(0, 0).text = "No"
            total_table.cell(0, 1).text = "Perincian Biaya"
            total_table.cell(0, 2).text = "Jumlah"
            total_table.cell(1, 0).merge(total_table.cell(1, 1)).text = "Jumlah"
            total_table.cell(1, 2).text = "Rp 730,000"
            receipt_table = document.add_table(rows=1, cols=1)
            receipt_table.cell(0, 0).text = "Rp 730,000"
            untouched_table = document.add_table(rows=1, cols=1)
            placeholder_table = document.add_table(rows=1, cols=1)
            thin_border_table = document.add_table(rows=1, cols=1)

            def apply_black_shading(cell):
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), "000000")
                cell._tc.get_or_add_tcPr().append(shading)

            def apply_black_end_border(cell, size):
                borders = OxmlElement("w:tcBorders")
                border = OxmlElement("w:end")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), str(size))
                border.set(qn("w:color"), "#000000")
                borders.append(border)
                cell._tc.get_or_add_tcPr().append(borders)

            apply_black_shading(total_table.cell(1, 0))
            apply_black_shading(total_table.cell(1, 2))
            apply_black_shading(receipt_table.cell(0, 0))
            apply_black_shading(untouched_table.cell(0, 0))
            apply_black_end_border(placeholder_table.cell(0, 0), 48)
            apply_black_end_border(thin_border_table.cell(0, 0), 8)
            document.save(path)

            repair_editable_docx(
                path,
                [
                    "RINCIAN BIAYA PERJALANAN DINAS\n"
                    "No Perincian Biaya Jumlah Keterangan"
                ],
                rincian_biaya_perjalanan_dinas=True,
            )

            result = Document(path)
            repaired_cells = (
                result.tables[0].cell(1, 0),
                result.tables[0].cell(1, 2),
                result.tables[1].cell(0, 0),
            )
            for cell in repaired_cells:
                shading = cell._tc.tcPr.find(qn("w:shd"))
                self.assertEqual(shading.get(qn("w:fill")), "E7E6E6")
                self.assertEqual(
                    cell.paragraphs[0].runs[0]._r.rPr.find(qn("w:color")).get(
                        qn("w:val")
                    ),
                    "000000",
                )

            untouched_shading = result.tables[2].cell(0, 0)._tc.tcPr.find(
                qn("w:shd")
            )
            self.assertEqual(untouched_shading.get(qn("w:fill")), "000000")
            placeholder_border = result.tables[3].cell(0, 0)._tc.tcPr.find(
                qn("w:tcBorders")
            ).find(qn("w:end"))
            self.assertEqual(placeholder_border.get(qn("w:color")), "#E7E6E6")
            thin_border = result.tables[4].cell(0, 0)._tc.tcPr.find(
                qn("w:tcBorders")
            ).find(qn("w:end"))
            self.assertEqual(thin_border.get(qn("w:color")), "#000000")

    def test_restores_form_title_underlines_without_retaining_spj_table_line(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "form-title-underlines.docx")
            document = Document()
            nota_title = document.add_paragraph()
            nota_title.add_run("\t")
            nota_run = nota_title.add_run("DAFTAR PENGELUARAN RIIL ")
            nota_body = nota_title.add_run("Yang bertanda tangan dibawah ini :")
            spj_title = document.add_paragraph()
            spj_run = spj_title.add_run("RINCIAN BIAYA PERJALANAN DINAS")
            field_table = document.add_table(rows=1, cols=3)
            field_table.cell(0, 0).text = "Lampiran SPD No."
            field_table.cell(0, 1).text = ":"
            field_table.cell(0, 2).text = "SPD-7/BPDP.33/2026"
            for cell in field_table.rows[0].cells:
                borders = OxmlElement("w:tcBorders")
                top = OxmlElement("w:top")
                top.set(qn("w:val"), "single")
                top.set(qn("w:sz"), "6")
                top.set(qn("w:color"), "#000000")
                borders.append(top)
                cell._tc.get_or_add_tcPr().append(borders)
            document.save(path)

            repair_editable_docx(
                path,
                [
                    "DAFTAR PENGELUARAN RIIL\nNo Uraian Jumlah\n"
                    "RINCIAN BIAYA PERJALANAN DINAS\n"
                    "Perincian Biaya Jumlah Keterangan"
                ],
                nota_riil=True,
                rincian_biaya_perjalanan_dinas=True,
            )

            result = Document(path)
            self.assertTrue(result.paragraphs[0].runs[1].font.underline)
            self.assertIsNone(result.paragraphs[0].runs[2].font.underline)
            self.assertTrue(result.paragraphs[1].runs[0].font.underline)
            for cell in result.tables[0].rows[0].cells:
                borders = cell._tc.tcPr.find(qn("w:tcBorders"))
                self.assertIsNone(borders.find(qn("w:top")))

    def test_restores_bpdp_header_rule_for_nota_and_spj_forms(self):
        cases = (
            (
                "nota-riil-header.docx",
                "DAFTAR PENGELUARAN RIIL",
                {"nota_riil": True},
            ),
            (
                "rincian-biaya-header.docx",
                "RINCIAN BIAYA PERJALANAN DINAS",
                {"rincian_biaya_perjalanan_dinas": True},
            ),
        )
        with tempfile.TemporaryDirectory() as directory:
            for filename, title, flags in cases:
                path = os.path.join(directory, filename)
                document = Document()
                document.add_paragraph(
                    "TELP (021) 84283099, SITUS www.bpdp.or.id"
                )
                document.add_paragraph(title)
                document.save(path)

                repair_editable_docx(path, [title], **flags)

                result = Document(path)
                rule = result.paragraphs[0]._p.getnext()
                self.assertEqual(rule.tag, qn("w:p"))
                properties = rule.find(qn("w:pPr"))
                bottom = properties.find(qn("w:pBdr")).find(qn("w:bottom"))
                self.assertEqual(bottom.get(qn("w:val")), "single")
                self.assertEqual(bottom.get(qn("w:sz")), "6")
                self.assertEqual(bottom.get(qn("w:color")), "000000")

    def test_restores_full_width_rincian_payment_band(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "rincian-payment-band.docx")
            document = Document()
            document.add_paragraph("RINCIAN BIAYA PERJALANAN DINAS")
            payment = document.add_table(rows=2, cols=2)
            nested_total = payment.cell(0, 0).add_table(rows=1, cols=1)
            nested_total.cell(0, 0).text = "Rp 730,000"
            payment.cell(0, 1).text = "Rp 730,000"
            payment.cell(1, 0).text = "Bendahara Pengeluaran"
            payment.cell(1, 1).text = "Yang menerima"
            document.save(path)

            repair_editable_docx(
                path,
                ["RINCIAN BIAYA PERJALANAN DINAS"],
                rincian_biaya_perjalanan_dinas=True,
            )

            result = Document(path)
            payment = result.tables[0]
            for cell in payment.rows[0].cells:
                shading = cell._tc.tcPr.find(qn("w:shd"))
                self.assertEqual(shading.get(qn("w:fill")), "E7E6E6")
            nested_total = payment.cell(0, 0).tables[0].cell(0, 0)
            shading = nested_total._tc.tcPr.find(qn("w:shd"))
            self.assertEqual(shading.get(qn("w:fill")), "E7E6E6")


if __name__ == "__main__":
    unittest.main()
